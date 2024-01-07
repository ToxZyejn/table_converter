import pandas as pd
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Border, Side

def set_cell_borders(cell):
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))
    cell.border = thin_border

def docx_to_excel(docx_path, excel_path):
    doc = Document(docx_path)
    workbook = Workbook()
    sheet = workbook.active

    row_offset = 1
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cell_value = cell.text
                excel_cell = sheet.cell(row=row_index + row_offset, column=col_index + 1)
                excel_cell.value = cell_value
                set_cell_borders(excel_cell)
        row_offset += len(table.rows) + 2  # Добавляем пустые строки между таблицами

    workbook.save(excel_path)

# Путь к вашему DOCX файлу
docx_path = 'extracted_tables.docx'
# Путь, куда сохранять Excel файл
excel_path = 'output_excel_file.xlsx'

docx_to_excel(docx_path, excel_path)
